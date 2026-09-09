import shutil
from datetime import date
from pathlib import Path

import openpyxl

from app.faktura_reader import KAT_JUZ_WYPELNIONE, KAT_SUKCES, KAT_ZUZYCIE_NIEZGODNE_Z_FAKTURA, PozycjaFakturyDuze
from app.import_faktur_duze import (
    polacz_pozycje_wielookresowe,
    wpisz_fakture_duzy_odbior_do_arkusza,
    wpisz_polaczone_pozycje_duze,
)

TESTY_FILE = Path(__file__).resolve().parents[2] / "examples" / "MPECWIK 2026_testy.xlsx"
FAKTURY_ROOT = Path(__file__).resolve().parents[2] / "examples" / "Faktury MPECWIK 2026"
FAKTURA_BRODOWO = next(p for p in FAKTURY_ROOT.rglob("D 03 31.pdf") if p.parent.name.startswith("Brodowo"))
FAKTURA_CHWALKOWO = next(p for p in FAKTURY_ROOT.rglob("D 01.31.pdf") if p.parent.name.startswith("Chwałkowo"))
FAKTURA_KORNICKA_STYCZEN = next(p for p in FAKTURY_ROOT.rglob("D 01 31.pdf") if p.parent.name.startswith("Kórnicka 82"))
FAKTURA_KORNICKA_LUTY_1 = next(p for p in FAKTURY_ROOT.rglob("D 02 19.pdf") if p.parent.name.startswith("Kórnicka 82"))
FAKTURA_KORNICKA_LUTY_2 = next(p for p in FAKTURY_ROOT.rglob("D 02 28.pdf") if p.parent.name.startswith("Kórnicka 82"))

# Hydrofornia Brodowo (Kod PPE 590310600000411983) to wiersz 4, blok marca = kolumny AU:BJ.
WIERSZ_BRODOWO = 4
# Oczyszczalnia ścieków Chwałkowo (taryfa B23, 3-strefowa) to wiersz 9, blok stycznia = kolumny M:AB.
WIERSZ_CHWALKOWO = 9
# Kórnicka 82 (taryfa B12, dzienna/nocna) to wiersz 7, blok stycznia M:AB, blok lutego AD:AS.
WIERSZ_KORNICKA = 7


def _kopia_testy(tmp_path) -> Path:
    cel = tmp_path / "testy.xlsx"
    shutil.copy(TESTY_FILE, cel)
    return cel


def test_wpisuje_mimo_niezgodnosci_zuzycia_i_zwraca_ostrzezenie(tmp_path):
    excel = _kopia_testy(tmp_path)

    wyniki = wpisz_fakture_duzy_odbior_do_arkusza(excel, FAKTURA_BRODOWO, nadpisuj=True)

    assert len(wyniki) == 1
    wynik = wyniki[0]
    assert wynik.kategoria == KAT_ZUZYCIE_NIEZGODNE_Z_FAKTURA
    assert wynik.wiersz == WIERSZ_BRODOWO
    assert wynik.nazwa_obiektu == "Hydrofornia Brodowo"
    assert "15818" in wynik.opis and "15793" in wynik.opis

    wb = openpyxl.load_workbook(excel, data_only=False)
    ws = wb["Duże odbiory"]
    assert ws["AU4"].value == 39
    assert ws["AV4"].value == 6807
    assert ws["AW4"].value == 9011
    assert ws["AY4"].value == 3470
    assert ws["AZ4"].value == 3216
    assert ws["BB4"].value == 50
    assert ws["BC4"].value == 211
    assert ws["BE4"].value == 447.62
    assert ws["BF4"].value == 406.23
    assert ws["BH4"].value == 5857.71
    # formuła "Razem zużycie" nie zostaje nadpisana literałem
    assert ws["BG4"].value == "=AV4+AW4+AX4"


def test_nie_nadpisuje_wypelnionych_komorek_bez_zgody(tmp_path):
    excel = _kopia_testy(tmp_path)
    # marzec dla Brodowo jest puste w danych testowych - wypełniamy najpierw sami (nadpisuj=True),
    # żeby drugie wpisanie (nadpisuj=False) miało na czym sprawdzić "już wypełnione".
    wpisz_fakture_duzy_odbior_do_arkusza(excel, FAKTURA_BRODOWO, nadpisuj=True)
    wyniki = wpisz_fakture_duzy_odbior_do_arkusza(excel, FAKTURA_BRODOWO, nadpisuj=False)

    assert len(wyniki) == 1
    assert wyniki[0].kategoria == KAT_JUZ_WYPELNIONE
    assert not wyniki[0].zapisano


def test_taryfa_3_strefowa_wpisuje_p_s3_i_q_s3_bez_ostrzezenia(tmp_path):
    # Chwałkowo: z doliczonymi licznikami strat (patrz faktura_reader.py) suma P-S1+P-S2+P-S3
    # zgadza się dokładnie z fakturą - żadnego KAT_ZUZYCIE_NIEZGODNE_Z_FAKTURA.
    excel = _kopia_testy(tmp_path)

    wyniki = wpisz_fakture_duzy_odbior_do_arkusza(excel, FAKTURA_CHWALKOWO, nadpisuj=True)

    assert len(wyniki) == 1
    assert wyniki[0].kategoria == KAT_SUKCES
    assert wyniki[0].wiersz == WIERSZ_CHWALKOWO

    wb = openpyxl.load_workbook(excel, data_only=False)
    ws = wb["Duże odbiory"]
    assert ws["M9"].value == 414  # Moc pobrana
    assert ws["N9"].value == 24387  # P-S1 (24383 odczyt + 4 strata)
    assert ws["O9"].value == 10853  # P-S2 (10850 odczyt + 3 strata)
    assert ws["P9"].value == 104245  # P-S3 (bez korekty - brak 3. licznika strat)
    assert ws["Q9"].value == 8233  # Q+ S1
    assert ws["R9"].value == 3164  # Q+ S2
    assert ws["S9"].value == 38280  # Q+ S3
    assert ws["T9"].value == 8  # Q- S1
    assert ws["U9"].value == 71  # Q- S2
    assert ws["V9"].value == 111  # Q- S3
    assert ws["Z9"].value == 29539.22  # Opłata netto dystrybucja


def test_taryfa_2_strefowa_nie_wpisuje_literalnego_zera_w_kolumny_s3(tmp_path):
    # Brodowo (2-strefowa) - P-S3/Q+S3/Q-S3 mają zostać puste, nie dostać literalne 0.
    excel = _kopia_testy(tmp_path)

    wpisz_fakture_duzy_odbior_do_arkusza(excel, FAKTURA_BRODOWO, nadpisuj=True)

    wb = openpyxl.load_workbook(excel, data_only=False)
    ws = wb["Duże odbiory"]
    assert ws["AX4"].value is None  # P-S3
    assert ws["BA4"].value is None  # Q+ S3
    assert ws["BD4"].value is None  # Q- S3


def test_niezgodnosc_nie_blokuje_zapisu_mimo_ze_zapisano_property_jest_false(tmp_path):
    # KAT_ZUZYCIE_NIEZGODNE_Z_FAKTURA to jedyna kategoria "z problemem", która MIMO TO zapisuje
    # dane (patrz import_faktur_duze.py) - inaczej niż wszystkie pozostałe kategorie problemów.
    # `zapisano` (WynikWpisu.zapisano, import_faktur.py) patrzy tylko na KAT_SUKCES, więc tu
    # celowo zostaje False, choć zapis realnie się odbył (potwierdzone wyżej wartościami komórek).
    excel = _kopia_testy(tmp_path)
    wyniki = wpisz_fakture_duzy_odbior_do_arkusza(excel, FAKTURA_BRODOWO, nadpisuj=True)

    assert wyniki[0].kategoria == KAT_ZUZYCIE_NIEZGODNE_Z_FAKTURA
    assert wyniki[0].zapisano is False


def test_taryfa_b12_dzienna_nocna_i_plus_procent_w_numerze_licznika(tmp_path):
    # Kórnicka 82: strefy "dzienna"/"nocna" (nie szczytowa/pozaszczytowa) i numer licznika z
    # doklejonym "+ 3%"/"+ 10%" w tekście faktury - oba potwierdzone jako osobne przyczyny, dla
    # których ta faktura się wcześniej nie parsowała w ogóle.
    excel = _kopia_testy(tmp_path)

    wyniki = wpisz_fakture_duzy_odbior_do_arkusza(excel, FAKTURA_KORNICKA_STYCZEN, nadpisuj=True)

    assert len(wyniki) == 1
    assert wyniki[0].kategoria == KAT_SUKCES
    assert wyniki[0].wiersz == WIERSZ_KORNICKA

    wb = openpyxl.load_workbook(excel, data_only=False)
    ws = wb["Duże odbiory"]
    assert ws["M7"].value == 12  # Moc pobrana
    assert ws["N7"].value == 4341  # P-S1 (dzienna)
    assert ws["O7"].value == 1681  # P-S2 (nocna)
    assert ws["Q7"].value == 2098  # Q+ S1
    assert ws["R7"].value == 743  # Q+ S2
    assert ws["T7"].value == 0  # Q- S1
    assert ws["U7"].value == 3  # Q- S2
    assert ws["X7"].value == 1.56  # Q- (zł)
    assert ws["Z7"].value == 2054.55  # Opłata netto dystrybucja


def test_kornicka_luty_dwie_faktury_sumuja_sie_w_jedna_pozycje(tmp_path):
    # Zmiana taryfy w środku lutego dała dwie osobne faktury za połówki miesiąca (1-19 i 20-28) -
    # wpisz_polaczone_pozycje_duze musi je zsumować w jedną pozycję, zamiast druga dostała "już
    # wypełnione" i zgubiła połowę miesiąca. Wartości zweryfikowane liczbowo wobec arkusza.
    excel = _kopia_testy(tmp_path)

    PPE_KORNICKA = "590310600000411990"
    wynik_per_plik, wlasne_okresy = wpisz_polaczone_pozycje_duze(
        excel, [FAKTURA_KORNICKA_LUTY_1, FAKTURA_KORNICKA_LUTY_2], nadpisuj=True
    )

    assert set(wynik_per_plik) == {FAKTURA_KORNICKA_LUTY_1, FAKTURA_KORNICKA_LUTY_2}
    # oba pliki dostają dokładnie ten sam (jeden) wpis - wspólnie złożyły się na tę samą pozycję
    assert wynik_per_plik[FAKTURA_KORNICKA_LUTY_1] == wynik_per_plik[FAKTURA_KORNICKA_LUTY_2]
    wynik = wynik_per_plik[FAKTURA_KORNICKA_LUTY_1][0]
    assert wynik.kategoria == KAT_SUKCES
    assert wynik.wiersz == WIERSZ_KORNICKA
    assert wynik.okres_od == date(2026, 2, 1)
    assert wynik.okres_do == date(2026, 2, 28)
    # ale KAŻDY plik zachowuje swój WŁASNY okres per PPE (do nazwy pliku w archiwum i do rozpoznania
    # "to poszło przez Duże odbiory" bez zgadywania po strukturze folderów - patrz foldery_obiektow.py)
    assert wlasne_okresy[FAKTURA_KORNICKA_LUTY_1] == {PPE_KORNICKA: date(2026, 2, 19)}
    assert wlasne_okresy[FAKTURA_KORNICKA_LUTY_2] == {PPE_KORNICKA: date(2026, 2, 28)}

    wb = openpyxl.load_workbook(excel, data_only=False)
    ws = wb["Duże odbiory"]
    assert ws["AD7"].value == 14  # Moc pobrana - MAKSIMUM (0 i 14), nie suma
    assert ws["AE7"].value == 3906  # P-S1 - suma (2711 + 1195)
    assert ws["AF7"].value == 1670  # P-S2 - suma (1192 + 478)
    assert ws["AH7"].value == 2489  # Q+ S1
    assert ws["AI7"].value == 985  # Q+ S2
    assert ws["AK7"].value == 1  # Q- S1
    assert ws["AL7"].value == 3  # Q- S2
    assert ws["AO7"].value == 2.08  # Q- (zł) - suma (1.56 + 0.52)
    assert round(ws["AQ7"].value, 2) == 2182.18  # Opłata netto dystrybucja - suma (1407.46 + 774.72)


def _pozycja(ppe="PPE", miesiac=2, moc=10.0, p_s1=100.0) -> PozycjaFakturyDuze:
    return PozycjaFakturyDuze(
        ppe=ppe,
        okres_od=date(2026, miesiac, 1),
        okres_do=date(2026, miesiac, 19),
        moc_pobrana=moc,
        p_s1=p_s1,
        p_s2=0.0,
        p_s3=0.0,
        q_plus_s1=0.0,
        q_plus_s2=0.0,
        q_plus_s3=0.0,
        q_minus_s1=0.0,
        q_minus_s2=0.0,
        q_minus_s3=0.0,
        q_plus_zl=0.0,
        q_minus_zl=0.0,
        oplata_dystrybucja=50.0,
        zuzycie_razem_faktura=100.0,
        trzy_strefy=False,
    )


def test_polacz_pozycje_wielookresowe_bez_kolizji_zostaje_bez_zmian():
    plik = Path("a.pdf")
    grupy = polacz_pozycje_wielookresowe([(plik, _pozycja())])

    assert grupy == [([plik], _pozycja())]


def test_polacz_pozycje_wielookresowe_sumuje_energie_ale_bierze_maksimum_mocy():
    plik1, plik2 = Path("a.pdf"), Path("b.pdf")
    p1 = _pozycja(moc=10.0, p_s1=100.0)
    p2 = _pozycja(moc=25.0, p_s1=50.0)

    grupy = polacz_pozycje_wielookresowe([(plik1, p1), (plik2, p2)])

    assert len(grupy) == 1
    pliki, polaczona = grupy[0]
    assert set(pliki) == {plik1, plik2}
    assert polaczona.moc_pobrana == 25.0  # maksimum, nie suma (35.0)
    assert polaczona.p_s1 == 150.0  # suma
    assert polaczona.oplata_dystrybucja == 100.0  # suma
    assert polaczona.okres_od == date(2026, 2, 1)
    assert polaczona.okres_do == date(2026, 2, 19)


def test_polacz_pozycje_wielookresowe_rozne_ppe_i_miesiace_zostaja_osobno():
    plik1, plik2, plik3 = Path("a.pdf"), Path("b.pdf"), Path("c.pdf")
    grupy = polacz_pozycje_wielookresowe(
        [
            (plik1, _pozycja(ppe="X", miesiac=2)),
            (plik2, _pozycja(ppe="Y", miesiac=2)),  # inne PPE - nie łączy się z X
            (plik3, _pozycja(ppe="X", miesiac=3)),  # ten sam PPE, inny miesiąc - też osobno
        ]
    )

    assert len(grupy) == 3
