import shutil
import zipfile
from pathlib import Path
from xml.etree import ElementTree

import openpyxl
import pytest

from app import xlsx_patch
from app.excel_reader import SHEET_NAME

TESTY_FILE = Path(__file__).resolve().parents[2] / "examples" / "MPECWIK 2026_testy.xlsx"


def _kopia(tmp_path) -> Path:
    cel = tmp_path / "kopia.xlsx"
    shutil.copy(TESTY_FILE, cel)
    return cel


def _kopia_z_calc_chain(tmp_path) -> Path:
    """Kopia pliku testowego z zagwarantowanym xl/calcChain.xml. Fixture obecnie go nie ma (pełny
    zapis podczas czyszczenia danych testowych Dużych odbiorów - openpyxl.save() gubi calcChain.xml,
    patrz komentarz modułu w xlsx_patch.py) - dopisujemy minimalny, ale poprawny wpis samodzielnie,
    żeby te dwa testy nie zależały od przypadkowego stanu fixture'u."""
    plik = _kopia(tmp_path)
    with zipfile.ZipFile(plik) as archiwum:
        if xlsx_patch.CALC_CHAIN in archiwum.namelist():
            return plik
        rels_xml = archiwum.read(xlsx_patch.WORKBOOK_RELS).decode("utf-8")
        content_types_xml = archiwum.read(xlsx_patch.CONTENT_TYPES).decode("utf-8")
        wpisy = {n: archiwum.read(n) for n in archiwum.namelist()}

    rels_xml = rels_xml.replace(
        "</Relationships>",
        '<Relationship Id="rIdCalcChainTest" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain" '
        'Target="calcChain.xml"/></Relationships>',
    )
    content_types_xml = content_types_xml.replace(
        "</Types>",
        '<Override PartName="/xl/calcChain.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml"/></Types>',
    )
    calc_chain_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<calcChain xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        '<c r="S3" i="1"/></calcChain>'
    )

    with zipfile.ZipFile(plik, "w", zipfile.ZIP_DEFLATED) as archiwum:
        for nazwa, tresc in wpisy.items():
            if nazwa == xlsx_patch.WORKBOOK_RELS:
                archiwum.writestr(nazwa, rels_xml)
            elif nazwa == xlsx_patch.CONTENT_TYPES:
                archiwum.writestr(nazwa, content_types_xml)
            else:
                archiwum.writestr(nazwa, tresc)
        archiwum.writestr(xlsx_patch.CALC_CHAIN, calc_chain_xml)
    return plik


def _nazwy_i_bajty(sciezka: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(sciezka) as archiwum:
        return {n: archiwum.read(n) for n in archiwum.namelist()}


def test_wpisuje_wartosci_i_mozna_je_odczytac_przez_openpyxl(tmp_path):
    plik = _kopia(tmp_path)
    xlsx_patch.wpisz_wartosci(plik, SHEET_NAME, {"P3": "01.01-20.01", "Q3": 182.0, "R3": 442.0, "T3": 244.41})

    wb = openpyxl.load_workbook(plik)
    ws = wb[SHEET_NAME]
    assert ws["P3"].value == "01.01-20.01"
    assert ws["Q3"].value == 182
    assert ws["R3"].value == 442
    assert ws["T3"].value == 244.41


def test_nie_rusza_formul_ani_innych_komorek(tmp_path):
    plik = _kopia(tmp_path)
    xlsx_patch.wpisz_wartosci(plik, SHEET_NAME, {"Q3": 182.0, "R3": 442.0})

    wb = openpyxl.load_workbook(plik)
    ws = wb[SHEET_NAME]
    assert ws["S3"].value == "=Q3+R3"
    assert ws["V3"].value == "=T3+U3"


def test_zachowuje_styl_komorki(tmp_path):
    plik = _kopia(tmp_path)
    with zipfile.ZipFile(plik) as archiwum:
        przed = archiwum.read("xl/worksheets/sheet1.xml").decode("utf-8")
    import re

    styl_przed = re.search(r'<c r="Q3" s="(\d+)"', przed).group(1)

    xlsx_patch.wpisz_wartosci(plik, SHEET_NAME, {"Q3": 182.0})

    with zipfile.ZipFile(plik) as archiwum:
        po = archiwum.read("xl/worksheets/sheet1.xml").decode("utf-8")
    styl_po = re.search(r'<c r="Q3" s="(\d+)"', po).group(1)
    assert styl_po == styl_przed


def test_zapisana_czesc_arkusza_jest_poprawnym_xmlem(tmp_path):
    plik = _kopia(tmp_path)
    xlsx_patch.wpisz_wartosci(plik, SHEET_NAME, {"P3": "01.01-20.01", "Q3": 182.0, "R3": 442.0, "T3": 244.41})
    with zipfile.ZipFile(plik) as archiwum:
        xml = archiwum.read("xl/worksheets/sheet1.xml").decode("utf-8")
    ElementTree.fromstring(xml)  # rzuci wyjątek, jeśli XML jest niepoprawny


def test_reszta_archiwum_zostaje_bajt_w_bajt_bez_zmian(tmp_path):
    """To jest właściwy test na problem, który naprawiliśmy: openpyxl.save() psuło
    externalLinks/sharedStrings/printerSettings przy pełnym zapisie skoroszytu, co Excel
    zgłaszał jako uszkodzony plik. Tu sprawdzamy, że poza samym arkuszem NIC innego się nie
    zmienia bajt w bajt."""
    plik = _kopia(tmp_path)
    przed = _nazwy_i_bajty(plik)

    xlsx_patch.wpisz_wartosci(plik, SHEET_NAME, {"Q3": 182.0, "R3": 442.0})

    po = _nazwy_i_bajty(plik)
    assert set(po) == set(przed)  # ta sama lista plików w archiwum
    for nazwa in przed:
        if nazwa == "xl/worksheets/sheet1.xml":
            continue
        assert po[nazwa] == przed[nazwa], f"{nazwa} zmieniło się, a nie powinno"


def test_wymus_przeliczenie_ustawia_fullcalconload_i_nie_dubluje_go(tmp_path):
    plik = _kopia(tmp_path)
    xlsx_patch.wymus_przeliczenie_formul(plik)
    with zipfile.ZipFile(plik) as archiwum:
        xml1 = archiwum.read("xl/workbook.xml").decode("utf-8")
    assert xml1.count('fullCalcOnLoad="1"') == 1

    xlsx_patch.wymus_przeliczenie_formul(plik)  # drugie wywolanie - idempotentne
    with zipfile.ZipFile(plik) as archiwum:
        xml2 = archiwum.read("xl/workbook.xml").decode("utf-8")
    assert xml2.count('fullCalcOnLoad="1"') == 1


def test_wymus_przeliczenie_nie_rusza_innych_czesci_archiwum(tmp_path):
    plik = _kopia(tmp_path)
    przed = _nazwy_i_bajty(plik)

    xlsx_patch.wymus_przeliczenie_formul(plik)

    po = _nazwy_i_bajty(plik)
    assert set(po) == set(przed)
    for nazwa in przed:
        if nazwa == "xl/workbook.xml":
            continue
        assert po[nazwa] == przed[nazwa], f"{nazwa} zmieniło się, a nie powinno"


def test_brakujaca_komorka_rzuca_czytelny_blad(tmp_path):
    plik = _kopia(tmp_path)
    with pytest.raises(xlsx_patch.KomorkaNieZnaleziona):
        xlsx_patch.wpisz_wartosci(plik, SHEET_NAME, {"ZZ9999": 1.0})


def test_nadpisanie_komorki_z_formula_usuwa_calc_chain(tmp_path):
    """Regresja: nadpisanie komórki, która miała formułę (np. ktoś ręcznie wpisał
    "=178.99+387.55" łącząc dwie faktury), zostawiało w calcChain.xml wpis wskazujący na
    formułę, której już nie ma w arkuszu - Excel zgłaszał to jako uszkodzoną zawartość."""
    plik = _kopia_z_calc_chain(tmp_path)
    with zipfile.ZipFile(plik) as archiwum:
        assert xlsx_patch.CALC_CHAIN in archiwum.namelist()
        przed = archiwum.read("xl/worksheets/sheet1.xml").decode("utf-8")
    assert '<c r="S3"' in przed and "<f>" in przed  # S3 = "=Q3+R3", patrz inne testy w tym pliku

    xlsx_patch.wpisz_wartosci(plik, SHEET_NAME, {"S3": 624.0})

    wb = openpyxl.load_workbook(plik)
    assert wb[SHEET_NAME]["S3"].value == 624

    with zipfile.ZipFile(plik) as archiwum:
        nazwy = archiwum.namelist()
        assert xlsx_patch.CALC_CHAIN not in nazwy
        rels_xml = archiwum.read(xlsx_patch.WORKBOOK_RELS).decode("utf-8")
        content_types_xml = archiwum.read(xlsx_patch.CONTENT_TYPES).decode("utf-8")
    assert "calcChain" not in rels_xml
    assert "calcChain" not in content_types_xml


def test_nadpisanie_zwyklej_komorki_nie_rusza_calc_chain(tmp_path):
    plik = _kopia_z_calc_chain(tmp_path)
    with zipfile.ZipFile(plik) as archiwum:
        przed = archiwum.read(xlsx_patch.CALC_CHAIN)

    xlsx_patch.wpisz_wartosci(plik, SHEET_NAME, {"Q3": 182.0, "R3": 442.0})

    with zipfile.ZipFile(plik) as archiwum:
        assert xlsx_patch.CALC_CHAIN in archiwum.namelist()
        po = archiwum.read(xlsx_patch.CALC_CHAIN)
    assert po == przed
