import shutil
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from app.excel_reader import SHEET_NAME, find_block_starts, find_identity_headers, find_name_column
from app.faktura_reader import (
    KAT_JUZ_WYPELNIONE,
    KAT_NIEZGODNOSC_SUMY,
    KAT_PPE_NIEZNALEZIONE,
    PozycjaFaktury,
)
from app.import_faktur import _bloki_po_miesiacu, _wiersze_po_ppe, _wpisz_pozycje, wpisz_fakture_do_arkusza

EXAMPLE_FILE = Path(__file__).resolve().parents[2] / "examples" / "MPECWIK 2026.xlsx"
TESTY_FILE = Path(__file__).resolve().parents[2] / "examples" / "MPECWIK 2026_testy.xlsx"
FAKTURY_ROOT = Path(__file__).resolve().parents[2] / "examples" / "Faktury MPECWIK 2026"
FAKTURA_20_PAZDZIERNIKA = next(p for p in FAKTURY_ROOT.rglob("D 01.pdf") if "Pa" in p.parent.name)


def _kopia_testy(tmp_path) -> Path:
    cel = tmp_path / "testy.xlsx"
    shutil.copy(TESTY_FILE, cel)
    return cel


def _pozycja(ppe, szczytowa=100.0, pozaszczytowa=50.0, razem=None, oplata=99.0, jednostrefowa=False):
    return PozycjaFaktury(
        ppe=ppe,
        okres_od=date(2026, 1, 1),
        okres_do=date(2026, 1, 20),
        zuzycie_szczytowa=szczytowa,
        zuzycie_pozaszczytowa=pozaszczytowa,
        zuzycie_razem=razem if razem is not None else szczytowa + pozaszczytowa,
        oplata_dystrybucja=oplata,
        jednostrefowa=jednostrefowa,
    )


def _srodowisko(ws):
    name_col = find_name_column(ws)
    kolumny_startowe = find_block_starts(ws)
    pierwszy_blok = min(kolumny_startowe)
    identity_headers = find_identity_headers(ws, pierwszy_blok)
    bloki = _bloki_po_miesiacu(ws, kolumny_startowe)
    wiersze = _wiersze_po_ppe(ws, name_col, identity_headers.get("Kod PPE"))
    return bloki, wiersze


def test_wpisuje_wszystkie_pozycje_faktury_do_pustych_komorek(tmp_path):
    excel = _kopia_testy(tmp_path)
    wyniki = wpisz_fakture_do_arkusza(excel, FAKTURA_20_PAZDZIERNIKA, nadpisuj=False)

    assert len(wyniki) == 5
    assert all(w.zapisano for w in wyniki)

    wb = openpyxl.load_workbook(excel, data_only=False)
    ws = wb[SHEET_NAME]
    # obiekt "20 Pazdziernika 40B" to wiersz 3, styczen zaczyna sie w kolumnie 16 (P)
    assert ws.cell(row=3, column=16).value == "01.01-20.01"
    assert ws.cell(row=3, column=17).value == 182.0
    assert ws.cell(row=3, column=18).value == 442.0
    assert ws.cell(row=3, column=20).value == 244.41
    # formula "Razem zuzycie" nie zostaje nadpisana literalem
    assert ws.cell(row=3, column=19).value == "=Q3+R3"


def test_nie_nadpisuje_wypelnionych_komorek_bez_zgody(tmp_path):
    excel = _kopia_testy(tmp_path)
    wpisz_fakture_do_arkusza(excel, FAKTURA_20_PAZDZIERNIKA, nadpisuj=False)
    # drugie uruchomienie na tych samych, juz wypelnionych komorkach
    wyniki = wpisz_fakture_do_arkusza(excel, FAKTURA_20_PAZDZIERNIKA, nadpisuj=False)

    assert all(not w.zapisano for w in wyniki)
    assert all(w.kategoria == KAT_JUZ_WYPELNIONE for w in wyniki)


def test_nadpisuje_wypelnione_komorki_gdy_dozwolone(tmp_path):
    excel = _kopia_testy(tmp_path)
    wpisz_fakture_do_arkusza(excel, FAKTURA_20_PAZDZIERNIKA, nadpisuj=False)
    wyniki = wpisz_fakture_do_arkusza(excel, FAKTURA_20_PAZDZIERNIKA, nadpisuj=True)

    assert all(w.zapisano for w in wyniki)


def test_ppe_spoza_arkusza_zostaje_pominiete():
    wb = openpyxl.load_workbook(EXAMPLE_FILE, data_only=False)
    ws = wb[SHEET_NAME]
    bloki, wiersze = _srodowisko(ws)

    wyniki, zapisy = _wpisz_pozycje(ws, [_pozycja(ppe="99999999999999999")], bloki, wiersze, nadpisuj=False)

    assert len(wyniki) == 1
    assert not wyniki[0].zapisano
    assert wyniki[0].kategoria == KAT_PPE_NIEZNALEZIONE
    assert not zapisy


def test_niezgodnosc_sumy_ps1_ps2_z_zuzyciem_na_fakturze_i_tak_zapisuje_z_ostrzezeniem(tmp_path):
    # Niezgodność sumy nie blokuje zapisu (z decyzji użytkownika, ten sam mechanizm co w Duże
    # odbiory) - dane trafiają do arkusza, ale wynik ma kategorię ostrzegawczą do weryfikacji.
    # Kopia _testy (nie prawdziwy EXAMPLE_FILE) - potrzebujemy pustych komórek dla tego PPE/miesiąca.
    excel = _kopia_testy(tmp_path)
    wb = openpyxl.load_workbook(excel, data_only=False)
    ws = wb[SHEET_NAME]
    bloki, wiersze = _srodowisko(ws)

    zla_pozycja = _pozycja(ppe="590310600000423740", szczytowa=100.0, pozaszczytowa=50.0, razem=999.0)
    wyniki, zapisy = _wpisz_pozycje(ws, [zla_pozycja], bloki, wiersze, nadpisuj=False)

    assert len(wyniki) == 1
    assert not wyniki[0].zapisano  # `zapisano` patrzy tylko na KAT_SUKCES, patrz import_faktur.py
    assert wyniki[0].kategoria == KAT_NIEZGODNOSC_SUMY
    assert zapisy  # ...ale dane faktycznie miały trafić do arkusza
    wiersz = wyniki[0].wiersz
    assert zapisy[f"Q{wiersz}"] == 100.0
    assert zapisy[f"R{wiersz}"] == 50.0


def test_licznik_jednostrefowy_wpisuje_tylko_p_s1_nie_p_s2(tmp_path):
    excel = _kopia_testy(tmp_path)
    wb = openpyxl.load_workbook(excel, data_only=False)
    ws = wb[SHEET_NAME]
    bloki, wiersze = _srodowisko(ws)

    pozycja = _pozycja(ppe="590310600000423740", szczytowa=9.0, pozaszczytowa=0.0, razem=9.0, jednostrefowa=True)
    wyniki, zapisy = _wpisz_pozycje(ws, [pozycja], bloki, wiersze, nadpisuj=False)

    assert wyniki[0].zapisano
    wiersz = wyniki[0].wiersz
    kolumny = bloki[1]
    assert zapisy[f"{openpyxl.utils.get_column_letter(kolumny['zuzycie_szczytowa'])}{wiersz}"] == 9.0
    assert f"{openpyxl.utils.get_column_letter(kolumny['zuzycie_pozaszczytowa'])}{wiersz}" not in zapisy


def test_faktura_ustawia_pelne_przeliczenie_formul_przy_wczytaniu(tmp_path):
    excel = _kopia_testy(tmp_path)
    wpisz_fakture_do_arkusza(excel, FAKTURA_20_PAZDZIERNIKA, nadpisuj=False)

    wb = openpyxl.load_workbook(excel)
    assert wb.calculation.fullCalcOnLoad is True
