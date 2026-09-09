from pathlib import Path

import openpyxl

from app.excel_reader_oddana import (
    FIRST_DATA_ROW_ODDANA,
    SHEET_NAME_ODDANA,
    find_block_starts_oddana,
    find_identity_headers_oddana,
    find_name_column_oddana,
    read_block_columns_oddana,
)

PLIK = Path(__file__).resolve().parents[2] / "examples" / "MPECWIK 2026.xlsx"


def _arkusz():
    wb = openpyxl.load_workbook(PLIK, data_only=False)
    return wb[SHEET_NAME_ODDANA]


def test_znajduje_12_blokow_miesiecznych_bez_rocznego_podsumowania():
    ws = _arkusz()
    starty = find_block_starts_oddana(ws)

    assert len(starty) == 12
    # J=10 (styczeń) ... BB=54 (grudzień) - kolumna BK=63 (roczne podsumowanie "P-S1") wykluczona.
    assert starty[0] == openpyxl.utils.column_index_from_string("J")
    assert starty[-1] == openpyxl.utils.column_index_from_string("BB")
    assert openpyxl.utils.column_index_from_string("BK") not in starty


def test_czyta_kolumny_pierwszego_bloku():
    ws = _arkusz()
    starty = find_block_starts_oddana(ws)
    kolumny = read_block_columns_oddana(ws, starty[0])

    assert kolumny == {
        "p_s1": openpyxl.utils.column_index_from_string("J"),
        "p_s2": openpyxl.utils.column_index_from_string("K"),
        "p_s3": openpyxl.utils.column_index_from_string("L"),
        "energia_wyprodukowana": openpyxl.utils.column_index_from_string("M"),
    }


def test_znajduje_kolumne_nazwy_i_ppe():
    ws = _arkusz()
    name_col = find_name_column_oddana(ws)
    assert name_col == 1

    naglowki = find_identity_headers_oddana(ws, find_block_starts_oddana(ws)[0])
    assert naglowki["Kod PPE"] == openpyxl.utils.column_index_from_string("F")


def test_wiersze_obiektow_maja_ppe_od_pierwszego_wiersza_danych():
    ws = _arkusz()
    name_col = find_name_column_oddana(ws)
    ppe_col = find_identity_headers_oddana(ws, find_block_starts_oddana(ws)[0])["Kod PPE"]

    assert ws.cell(row=FIRST_DATA_ROW_ODDANA, column=name_col).value.strip() == "Babin"
    assert ws.cell(row=FIRST_DATA_ROW_ODDANA, column=ppe_col).value == "590310600000411976"
