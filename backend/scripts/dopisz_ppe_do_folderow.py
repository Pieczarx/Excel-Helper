"""Jednorazowy (ale idempotentny - bezpieczny do ponownego uruchomienia) skrypt: dopisuje Kod PPE
do nazw istniejących folderów obiektów w drzewie faktur (`<nazwa> (<PPE>)`), żeby runtime
(app/foldery_obiektow.py) mógł rozdzielać wpisane faktury po PPE zamiast dopasowania po nazwie.

Zmienia WYŁĄCZNIE nazwy katalogów - nigdy nazw plików faktur (np. "D 02.pdf" zostaje nietknięte).

Pomija foldery, które już mają `(<ppe>)` na końcu nazwy - drugie uruchomienie nic nie zmienia.

DOMYŚLNIE URUCHAMIA SIĘ NA SUCHO (nic nie zmienia na dysku) - wypisuje tylko, co by zrobił.
Żeby faktycznie zmienić nazwy, dodaj flagę --wykonaj na końcu.

Wymaga tego samego środowiska co aplikacja (openpyxl, pymupdf zainstalowane) - poza tym
wystarczy zwykłe `python`, nie trzeba nic ustawiać ręcznie (skrypt sam dokłada swój katalog
`backend/` do ścieżki importów):

    python scripts/dopisz_ppe_do_folderow.py "<ścieżka do drzewa faktur>" "<ścieżka do pliku Excel>"
    # po sprawdzeniu wypisanego planu:
    python scripts/dopisz_ppe_do_folderow.py "<ścieżka do drzewa faktur>" "<ścieżka do pliku Excel>" --wykonaj

Przykład:
    python scripts/dopisz_ppe_do_folderow.py "C:\\Faktury MPECWIK 2026" "C:\\MPECWIK 2026.xlsx" --wykonaj
"""
from __future__ import annotations

import re
import sys
import unicodedata
from difflib import SequenceMatcher
from pathlib import Path

# Żeby `python scripts\dopisz_ppe_do_folderow.py ...` działało z dowolnego katalogu roboczego
# i bez ręcznego ustawiania PYTHONPATH (np. w PowerShell samo "set PYTHONPATH=." nie ustawia
# zmiennej środowiskowej tak jak w cmd.exe - to inny alias, więc import 'app' i tak by się wywalił).
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl

from app.faktura_reader import NieRozpoznanoFaktury, wczytaj_fakture_duzy_odbior, wczytaj_fakture_dystrybucyjna

WYKLUCZ = {"Do wpisania", "Do aktualizacji", "Faktury powyżej 16"}
PROG_PEWNOSCI = 0.4

# Foldery, dla których PPE zostało już ustalone/potwierdzone ręcznie - ogólny algorytm dopasowania
# nazw albo w ogóle nie znajdzie tu PPE (faktura się nie parsuje), albo trafi z bardzo niskim
# wynikiem podobieństwa nazw (nazwa folderu nie przypomina nazwy w Excelu). Specyficzne dla
# naszego drzewa faktur - jeśli danego folderu tam nie ma, po prostu nic się nie dzieje.
ZNANE_PPE_DLA_FOLDEROW: dict[Path, str] = {
    # literówka w nazwie ("Kosynierórw") + faktura jednostrefowa "całodobowa", której duży-odbiorowy
    # parser jeszcze nie obsługuje - PPE ustalone przez wykluczenie (jedyny niedopasowany PPE
    # w Duże odbiory) i potwierdzone ręcznie.
    Path("DUŻE ODBIORY") / "Kosynierórw": "590310600000414366",
    # nazwa folderu (potoczna nazwa miejscowości) nie przypomina adresu ulicowego w Excelu -
    # potwierdzone ręcznie jako poprawne dopasowanie mimo niskiego podobieństwa nazw.
    Path("Włostowo"): "590310600000423689",
    # plik w środku nazywa się "DE 02.pdf" (literówka - nie dotykamy nazw plików, tylko katalogów),
    # więc ogólny algorytm może nie znaleźć/nie dopasować reprezentanta - PPE potwierdzone ręcznie.
    Path("Zielona"): "590310600032555914",
}

_WZORZEC_PPE_W_NAZWIE = re.compile(r"\((\d+)\)$")


def norm(s: str) -> str:
    s = unicodedata.normalize("NFKD", s.lower())
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def wczytaj_ppe_z_arkusza(sciezka_excel: Path) -> dict[str, tuple[str, str, int]]:
    """Kod PPE -> (nazwa arkusza, nazwa obiektu, numer wiersza)."""
    wb = openpyxl.load_workbook(sciezka_excel, data_only=False)
    ppe_do_nazwy: dict[str, tuple[str, str, int]] = {}
    for nazwa_arkusza in ("Małe odbiory", "Duże odbiory"):
        ws = wb[nazwa_arkusza]
        header_row = 2 if nazwa_arkusza == "Małe odbiory" else 1
        ppe_col = name_col = None
        for col in range(1, 30):
            v = ws.cell(row=header_row, column=col).value
            if v == "Kod PPE":
                ppe_col = col
            if v == "Ulica/ Nazwa własna":
                name_col = col
        for row in range(header_row + 1, ws.max_row + 1):
            nazwa = ws.cell(row=row, column=name_col).value
            ppe = ws.cell(row=row, column=ppe_col).value
            if nazwa and ppe:
                ppe_do_nazwy[str(ppe).strip()] = (nazwa_arkusza, str(nazwa).strip(), row)
    return ppe_do_nazwy


def znajdz_liscie(root: Path):
    for p in sorted(root.rglob("*")):
        if p.is_dir():
            if any(cz in WYKLUCZ for cz in p.relative_to(root).parts):
                continue
            pliki = sorted(p.glob("*.pdf"))
            if pliki:
                yield p, pliki


def wybierz_reprezentanta(pliki: list[Path]) -> Path:
    for p in pliki:
        if p.name.upper().startswith("D "):
            return p
    for p in pliki:
        if p.name.upper().startswith("D"):
            return p
    return pliki[0]


def sprobuj_ppe(sciezka: Path) -> set[str]:
    """PPE z w pełni rozpoznanych pozycji, a jeśli takich brak - też z sekcji pominiętych (mają
    PPE ustalone z samego nagłówka sekcji, nawet gdy reszta pól się nie sparsowała - wystarczy do
    dopasowania folderu, nie trzeba pełnego odczytu faktury)."""
    znalezione: set[str] = set()
    pominiete_ppe: set[str] = set()
    for wczytaj in (wczytaj_fakture_dystrybucyjna, wczytaj_fakture_duzy_odbior):
        try:
            wynik = wczytaj(sciezka)
            znalezione.update(p.ppe for p in wynik.pozycje)
            pominiete_ppe.update(s.ppe for s in wynik.pominiete if s.ppe)
        except NieRozpoznanoFaktury:
            pass
        except Exception:
            pass
    return znalezione or pominiete_ppe


def najlepsze_dopasowanie(folder_rel: Path, kandydaci: set[str], ppe_do_nazwy: dict) -> tuple | None:
    cel = " ".join(norm(cz) for cz in folder_rel.parts)
    najlepszy = None
    for ppe in kandydaci:
        info = ppe_do_nazwy.get(ppe)
        if info is None:
            continue
        nazwa_excel = norm(info[1])
        score = SequenceMatcher(None, cel, nazwa_excel).ratio()
        if nazwa_excel in cel or cel in nazwa_excel:
            score += 0.3
        if najlepszy is None or score > najlepszy[2]:
            najlepszy = (ppe, info, score)
    return najlepszy


def main() -> None:
    argumenty = [a for a in sys.argv[1:] if a != "--wykonaj"]
    na_sucho = "--wykonaj" not in sys.argv
    if len(argumenty) != 2:
        print(__doc__)
        raise SystemExit(1)
    root = Path(argumenty[0]).resolve()
    sciezka_excel = Path(argumenty[1]).resolve()

    if not root.is_dir():
        print(f"Nie znaleziono folderu: {root}")
        raise SystemExit(1)
    if not sciezka_excel.is_file():
        print(f"Nie znaleziono pliku Excel: {sciezka_excel}")
        raise SystemExit(1)

    print(f"Drzewo faktur: {root}")
    print(f"Plik Excel:    {sciezka_excel}")
    print("TRYB: na sucho (nic nie zmieniam) - dodaj --wykonaj, żeby faktycznie zmienić nazwy." if na_sucho else "TRYB: WYKONUJĘ ZMIANY")
    print()

    ppe_do_nazwy = wczytaj_ppe_z_arkusza(sciezka_excel)
    zmienione = 0
    pominiete: list[Path] = []
    bez_dopasowania: list[Path] = []

    obsluzone_przez_znane: set[Path] = set()  # w trybie na sucho dysk się nie zmienia, więc
    # rozpoznanie "już obsłużone" dla poniższej ogólnej pętli nie może polegać na stanie dysku.
    for stara_wzgl, ppe in ZNANE_PPE_DLA_FOLDEROW.items():
        stara = root / stara_wzgl
        if not stara.is_dir() or _WZORZEC_PPE_W_NAZWIE.search(stara.name):
            continue
        info = ppe_do_nazwy.get(ppe)
        nowa_nazwa = f"{stara.name} ({ppe})"
        print(f"ZMIENIONO (znane): {stara_wzgl} -> {nowa_nazwa}  ('{info[1] if info else '?'}')")
        if not na_sucho:
            stara.rename(stara.with_name(nowa_nazwa))
        obsluzone_przez_znane.add(stara)
        zmienione += 1

    for folder, pliki in znajdz_liscie(root):
        if folder in obsluzone_przez_znane:
            continue  # już wypisane wyżej jako "ZMIENIONO (znane)"
        if _WZORZEC_PPE_W_NAZWIE.search(folder.name):
            pominiete.append(folder.relative_to(root))
            continue
        rel = folder.relative_to(root)
        reprezentant = wybierz_reprezentanta(pliki)
        kandydaci = sprobuj_ppe(reprezentant)
        wynik = najlepsze_dopasowanie(rel, kandydaci, ppe_do_nazwy) if kandydaci else None
        if wynik is None or wynik[2] < PROG_PEWNOSCI:
            bez_dopasowania.append(rel)
            continue
        ppe, info, score = wynik
        nowa_nazwa = f"{folder.name} ({ppe})"
        print(f"ZMIENIONO: {rel} -> {nowa_nazwa}  (~{score:.2f}, '{info[1]}')")
        if not na_sucho:
            folder.rename(folder.with_name(nowa_nazwa))
        zmienione += 1

    print()
    print(f"{'Do zmiany' if na_sucho else 'Zmienionych'} folderów: {zmienione}")
    print(f"Już oznaczonych wcześniej ({len(pominiete)}):")
    for rel in pominiete:
        print(f"  {rel}")
    print(f"Bez dopasowania ({len(bez_dopasowania)}):")
    for rel in bez_dopasowania:
        print(f"  {rel}")

    if na_sucho and zmienione:
        print()
        print("To był przebieg na sucho - nic nie zostało zmienione na dysku.")
        print("Sprawdź powyższą listę, a potem uruchom ponownie z --wykonaj, żeby faktycznie zmienić nazwy.")


if __name__ == "__main__":
    main()
