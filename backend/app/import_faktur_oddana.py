"""Wpisuje pozycje energii oddanej (fotowoltaika) z faktur dystrybucyjnych PDF do arkusza
'fotowoltaika energia oddana', dopasowując obiekty po Kodzie PPE - mirror import_faktur_duze.py,
patrz tam po ogólny opis mechanizmu zapisu (xlsx_patch.py, formuły niedotykane, "nie nadpisuj bez
zgody"). Kolumna "Energia wyprodukowana" jest polem ręcznym użytkownika (patrz faktura_reader.py)
- ten moduł pisze tylko P-S1/P-S2/P-S3."""
from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

from app import xlsx_patch
from app.excel_reader import norm
from app.excel_reader_oddana import (
    FIRST_DATA_ROW_ODDANA,
    SHEET_NAME_ODDANA,
    find_block_starts_oddana,
    find_identity_headers_oddana,
    find_name_column_oddana,
    read_block_columns_oddana,
)
from app.faktura_reader import (
    KAT_JUZ_WYPELNIONE,
    KAT_PPE_NIEZNALEZIONE,
    KAT_SUKCES,
    NieRozpoznanoFaktury,
    PozycjaEnergiaOddana,
    wczytaj_fakture_energia_oddana,
)
from app.import_faktur import WynikWpisu

POLA_DO_ZAPISU_ODDANA = ("p_s1", "p_s2", "p_s3")


def _bloki_po_miesiacu_oddana(ws, kolumny_startowe: list[int]) -> dict[int, dict[str, int]]:
    """Numer miesiąca (1-12) -> mapowanie pole->kolumna. Arkusz nie ma wiersza z nazwami miesięcy
    - bloki są w kolejności występowania (1. = styczeń, ..., 12. = grudzień), jak w Duże odbiory."""
    return {indeks + 1: read_block_columns_oddana(ws, start) for indeks, start in enumerate(kolumny_startowe)}


def _wiersze_po_ppe_oddana(ws, name_col: int, ppe_col: int | None) -> dict[str, tuple[int, str]]:
    """Kod PPE -> (numer wiersza, nazwa obiektu) - lokalizacja obiektu w arkuszu ZAWSZE po PPE,
    nigdy po nazwie/pozycji folderu (ten sam mechanizm co w Małe/Duże odbiory)."""
    if ppe_col is None:
        return {}
    wynik: dict[str, tuple[int, str]] = {}
    for row in range(FIRST_DATA_ROW_ODDANA, ws.max_row + 1):
        nazwa = ws.cell(row=row, column=name_col).value
        if nazwa is None or norm(nazwa) == "":
            continue
        ppe = norm(ws.cell(row=row, column=ppe_col).value)
        if ppe:
            wynik[ppe] = (row, norm(nazwa))
    return wynik


def _komorki_puste_oddana(ws, wiersz: int, kolumny: dict[str, int]) -> bool:
    for pole in ("p_s1", "p_s2"):
        col = kolumny.get(pole)
        if col is not None and norm(ws.cell(row=wiersz, column=col).value) != "":
            return False
    return True


def _referencja(wiersz: int, kolumna: int) -> str:
    return f"{get_column_letter(kolumna)}{wiersz}"


def _wpisz_pozycje_oddana(
    ws,
    pozycje: list[PozycjaEnergiaOddana],
    bloki_po_miesiacu: dict[int, dict[str, int]],
    wiersze_po_ppe: dict[str, tuple[int, str]],
    nadpisuj: bool,
) -> tuple[list[WynikWpisu], dict[str, str | float]]:
    wyniki: list[WynikWpisu] = []
    zapisy: dict[str, str | float] = {}
    for pozycja in pozycje:
        miesiac = pozycja.okres_do.month
        trafienie = wiersze_po_ppe.get(pozycja.ppe)
        if trafienie is None:
            wyniki.append(
                WynikWpisu(
                    kategoria=KAT_PPE_NIEZNALEZIONE,
                    ppe=pozycja.ppe,
                    wiersz=None,
                    nazwa_obiektu=None,
                    miesiac=miesiac,
                    okres_od=pozycja.okres_od,
                    okres_do=pozycja.okres_do,
                    opis="brak obiektu z tym numerem PPE w arkuszu 'fotowoltaika energia oddana'",
                    arkusz=SHEET_NAME_ODDANA,
                )
            )
            continue
        wiersz, nazwa_obiektu = trafienie

        kolumny = bloki_po_miesiacu.get(miesiac)
        if kolumny is None:
            continue  # arkusz nie ma bloku dla tego miesiąca - nie powinno się zdarzyć, 12 bloków zawsze obecnych

        if not nadpisuj and not _komorki_puste_oddana(ws, wiersz, kolumny):
            wyniki.append(
                WynikWpisu(
                    kategoria=KAT_JUZ_WYPELNIONE,
                    ppe=pozycja.ppe,
                    wiersz=wiersz,
                    nazwa_obiektu=nazwa_obiektu,
                    miesiac=miesiac,
                    okres_od=pozycja.okres_od,
                    okres_do=pozycja.okres_do,
                    opis="komórki już wypełnione (tryb bez nadpisywania)",
                    arkusz=SHEET_NAME_ODDANA,
                )
            )
            continue

        wartosci = {"p_s1": pozycja.p_s1, "p_s2": pozycja.p_s2, "p_s3": pozycja.p_s3}
        # Taryfa 2-strefowa: p_s3 to fikcyjne 0.0 (patrz PozycjaEnergiaOddana.trzy_strefy) -
        # kolumna P-S3 ma zostać puste, nie dostać literalne zero (jak w Duże odbiory).
        pola = POLA_DO_ZAPISU_ODDANA if pozycja.trzy_strefy else ("p_s1", "p_s2")
        for pole in pola:
            col = kolumny.get(pole)
            if col is not None:
                zapisy[_referencja(wiersz, col)] = wartosci[pole]

        # Do pokazania w UI (WierszObiektu) - bez "Razem"/"Dystrybucja", ten arkusz ich nie ma
        # (patrz styl.py/widok_uzupelnij_excel.py po decyzję użytkownika).
        wartosci_wyswietlane = {"P-S1": pozycja.p_s1, "P-S2": pozycja.p_s2}
        if pozycja.trzy_strefy:
            wartosci_wyswietlane["P-S3"] = pozycja.p_s3

        wyniki.append(
            WynikWpisu(
                kategoria=KAT_SUKCES,
                ppe=pozycja.ppe,
                wiersz=wiersz,
                nazwa_obiektu=nazwa_obiektu,
                miesiac=miesiac,
                okres_od=pozycja.okres_od,
                okres_do=pozycja.okres_do,
                arkusz=SHEET_NAME_ODDANA,
                wartosci=wartosci_wyswietlane,
            )
        )
    return wyniki, zapisy


def wpisz_fakture_oddana_do_arkusza(
    sciezka_excel: str | Path, sciezka_faktury: str | Path, nadpisuj: bool
) -> list[WynikWpisu]:
    """Parsuje jedną fakturę dystrybucyjną PDF i wpisuje jej pozycje energii oddanej do arkusza
    'fotowoltaika energia oddana', dopasowując po Kodzie PPE. Zapisuje plik Excela na miejscu."""
    wynik_faktury = wczytaj_fakture_energia_oddana(sciezka_faktury)

    wb = openpyxl.load_workbook(sciezka_excel, data_only=False)
    ws = wb[SHEET_NAME_ODDANA]

    name_col = find_name_column_oddana(ws)
    kolumny_startowe = find_block_starts_oddana(ws)
    pierwszy_blok = min(kolumny_startowe) if kolumny_startowe else ws.max_column + 1
    identity_headers = find_identity_headers_oddana(ws, pierwszy_blok)

    bloki_po_miesiacu = _bloki_po_miesiacu_oddana(ws, kolumny_startowe)
    wiersze_po_ppe = _wiersze_po_ppe_oddana(ws, name_col, identity_headers.get("Kod PPE"))

    wyniki, zapisy = _wpisz_pozycje_oddana(ws, wynik_faktury.pozycje, bloki_po_miesiacu, wiersze_po_ppe, nadpisuj)

    if zapisy:
        xlsx_patch.wpisz_wartosci(sciezka_excel, SHEET_NAME_ODDANA, zapisy)
        xlsx_patch.wymus_przeliczenie_formul(sciezka_excel)

    return wyniki


def polacz_pozycje_wielookresowe_oddana(
    pozycje_ze_zrodlem: list[tuple[Path, PozycjaEnergiaOddana]],
) -> list[tuple[list[Path], PozycjaEnergiaOddana]]:
    """Łączy pozycje tego samego PPE i tego samego miesiąca docelowego z RÓŻNYCH plików w jedną,
    zsumowaną pozycję - ten sam mechanizm i to samo uzasadnienie co
    import_faktur_duze.polacz_pozycje_wielookresowe (zmiana taryfy w środku miesiąca -> dwie
    faktury za połówki miesiąca). P-S1/P-S2/P-S3 są addytywne (kWh za kawałek okresu), więc się
    sumują."""
    grupy: dict[tuple[str, int], list[tuple[Path, PozycjaEnergiaOddana]]] = {}
    kolejnosc: list[tuple[str, int]] = []
    for plik, pozycja in pozycje_ze_zrodlem:
        klucz = (pozycja.ppe, pozycja.okres_do.month)
        if klucz not in grupy:
            grupy[klucz] = []
            kolejnosc.append(klucz)
        grupy[klucz].append((plik, pozycja))

    polaczone: list[tuple[list[Path], PozycjaEnergiaOddana]] = []
    for klucz in kolejnosc:
        czlonkowie = grupy[klucz]
        pliki = [plik for plik, _ in czlonkowie]
        if len(czlonkowie) == 1:
            polaczone.append((pliki, czlonkowie[0][1]))
            continue

        pozycje = [pozycja for _, pozycja in czlonkowie]
        polaczone.append(
            (
                pliki,
                PozycjaEnergiaOddana(
                    ppe=pozycje[0].ppe,
                    okres_od=min(p.okres_od for p in pozycje),
                    okres_do=max(p.okres_do for p in pozycje),
                    p_s1=sum(p.p_s1 for p in pozycje),
                    p_s2=sum(p.p_s2 for p in pozycje),
                    p_s3=sum(p.p_s3 for p in pozycje),
                    trzy_strefy=pozycje[0].trzy_strefy,
                ),
            )
        )
    return polaczone


def wpisz_polaczone_pozycje_oddana(
    sciezka_excel: str | Path, sciezki_faktur: list[str | Path], nadpisuj: bool
) -> dict[Path, list[WynikWpisu]]:
    """Parsuje WSZYSTKIE podane faktury naraz (bez zapisu), łączy pozycje tego samego PPE/miesiąca
    rozbite na kilka plików i zapisuje połączony wynik do arkusza 'fotowoltaika energia oddana'
    jednym zapisem - mirror import_faktur_duze.wpisz_polaczone_pozycje_duze (bez drugiej mapy
    "własny okres_do": routing do folderów obiektów już obsługuje ten sam PPE przez przebieg Duże
    odbiory - fotowoltaika śledzi wyłącznie obiekty, które i tak przez niego przechodzą)."""
    sciezki = [Path(p) for p in sciezki_faktur]
    pozycje_ze_zrodlem: list[tuple[Path, PozycjaEnergiaOddana]] = []
    for sciezka in sciezki:
        try:
            wynik_faktury = wczytaj_fakture_energia_oddana(sciezka)
        except NieRozpoznanoFaktury:
            continue  # w ogóle nie faktura dystrybucyjna - zgłosi to osobno przebieg Małe/Duże odbiory
        pozycje_ze_zrodlem.extend((sciezka, p) for p in wynik_faktury.pozycje)

    grupy = polacz_pozycje_wielookresowe_oddana(pozycje_ze_zrodlem)

    wb = openpyxl.load_workbook(sciezka_excel, data_only=False)
    ws = wb[SHEET_NAME_ODDANA]
    name_col = find_name_column_oddana(ws)
    kolumny_startowe = find_block_starts_oddana(ws)
    pierwszy_blok = min(kolumny_startowe) if kolumny_startowe else ws.max_column + 1
    identity_headers = find_identity_headers_oddana(ws, pierwszy_blok)
    bloki_po_miesiacu = _bloki_po_miesiacu_oddana(ws, kolumny_startowe)
    wiersze_po_ppe = _wiersze_po_ppe_oddana(ws, name_col, identity_headers.get("Kod PPE"))

    pozycje_polaczone = [pozycja for _, pozycja in grupy]
    wyniki, zapisy = _wpisz_pozycje_oddana(ws, pozycje_polaczone, bloki_po_miesiacu, wiersze_po_ppe, nadpisuj)

    wynik_per_plik: dict[Path, list[WynikWpisu]] = {plik: [] for plik in sciezki}
    for (pliki_grupy, _pozycja), wpis in zip(grupy, wyniki):
        for plik in pliki_grupy:
            wynik_per_plik[plik].append(wpis)

    if zapisy:
        xlsx_patch.wpisz_wartosci(sciezka_excel, SHEET_NAME_ODDANA, zapisy)
        xlsx_patch.wymus_przeliczenie_formul(sciezka_excel)

    return wynik_per_plik
