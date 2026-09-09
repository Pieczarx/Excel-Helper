"""Odczyt struktury arkusza 'fotowoltaika energia oddana' - mirror excel_reader_duze.py, z jedną
realną różnicą w rozpoznawaniu bloków miesięcznych:

Blok miesięczny ma 4 kolumny (P-S1, P-S2, P-S3, Energia wyprodukowana) i zaczyna się nagłówkiem
'P-S1' w wierszu 1 - ALE ten sam nagłówek 'P-S1' występuje też w kolumnie BK, jako początek
rocznego podsumowania (P-S1/P-S2/Energia wyprodukowana/Razem, bez P-S3 w środku). Żeby nie wziąć
tego podsumowania za 13. "miesiąc", blok miesięczny musi mieć 'P-S3' dokładnie 2 kolumny dalej -
podsumowanie tego nie ma (P-S2 -> od razu Energia wyprodukowana), więc odpada.

Kolumna 'Energia wyprodukowana' w każdym bloku jest polem ręcznym użytkownika (brak odpowiednika
na jakiejkolwiek sprawdzonej fakturze D/E) - z decyzji użytkownika ten moduł/import_faktur_oddana.py
jej nie dotyka, tylko czyta P-S1/P-S2/P-S3."""
from __future__ import annotations

import openpyxl

from app.excel_reader import norm

SHEET_NAME_ODDANA = "fotowoltaika energia oddana"
HEADER_ROW_ODDANA = 1
FIRST_DATA_ROW_ODDANA = 2
NAME_COLUMN_HEADER_ODDANA = "Ulica/ Nazwa własna"
BLOCK_START_HEADER_ODDANA = "P-S1"
BLOCK_WIDTH_ODDANA = 4  # P-S1, P-S2, P-S3, Energia wyprodukowana

FIELD_ALIASES_ODDANA = {
    "P-S1": "p_s1",
    "P-S2": "p_s2",
    "P-S3": "p_s3",
    "Energia wyprodukowana": "energia_wyprodukowana",
}


def find_block_starts_oddana(ws) -> list[int]:
    """Kolumny, w których zaczyna się blok miesięczny ('P-S1' w wierszu 1, z 'P-S3' 2 kolumny dalej
    - patrz docstring modułu po co), z pominięciem ukrytych kolumn, w kolejności = styczeń..grudzień."""
    starts = []
    for col in range(1, ws.max_column + 1):
        letter = openpyxl.utils.get_column_letter(col)
        dim = ws.column_dimensions.get(letter)
        if dim is not None and dim.hidden:
            continue
        if norm(ws.cell(row=HEADER_ROW_ODDANA, column=col).value) != BLOCK_START_HEADER_ODDANA:
            continue
        if norm(ws.cell(row=HEADER_ROW_ODDANA, column=col + 2).value) != "P-S3":
            continue  # roczne podsumowanie (P-S1 w kolumnie BK) - nie blok miesięczny
        starts.append(col)
    return starts


def read_block_columns_oddana(ws, start_col: int) -> dict[str, int]:
    """Mapuje nazwę pola -> numer kolumny dla jednego bloku miesięcznego (4 kolumny), zaczynając
    od kolumny 'P-S1'."""
    columns: dict[str, int] = {}
    for col in range(start_col, start_col + BLOCK_WIDTH_ODDANA):
        header = norm(ws.cell(row=HEADER_ROW_ODDANA, column=col).value)
        alias = FIELD_ALIASES_ODDANA.get(header)
        if alias:
            columns[alias] = col
    return columns


def find_name_column_oddana(ws) -> int:
    for col in range(1, ws.max_column + 1):
        if norm(ws.cell(row=HEADER_ROW_ODDANA, column=col).value) == NAME_COLUMN_HEADER_ODDANA:
            return col
    raise ValueError(f"Nie znaleziono kolumny nagłówka {NAME_COLUMN_HEADER_ODDANA!r} w wierszu {HEADER_ROW_ODDANA}")


def find_identity_headers_oddana(ws, ostatnia_kolumna: int) -> dict[str, int]:
    """Nagłówek (znormalizowany) -> numer kolumny, dla kolumn tożsamości obiektu (przed 1. blokiem
    miesięcznym) - niezależnie od tego, czy kolumna jest ukryta."""
    naglowki: dict[str, int] = {}
    for col in range(1, ostatnia_kolumna):
        header = norm(ws.cell(row=HEADER_ROW_ODDANA, column=col).value)
        if header:
            naglowki[header] = col
    return naglowki
