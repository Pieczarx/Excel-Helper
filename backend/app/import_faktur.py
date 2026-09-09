"""Wpisuje pozycje z faktur dystrybucyjnych PDF do arkusza 'Małe odbiory', dopasowując obiekty
po Kodzie PPE. Faktura jest zapisywana pod blokiem miesiąca daty końcowej okresu (tak samo jak
przy odczycie - patrz excel_reader.py), a "Razem zużycie"/"Razem koszty" to formuły w arkuszu,
więc wpisujemy tylko Data/P-S1/P-S2/Opłata netto dystrybucja i pozwalamy formule policzyć resztę.

Sam zapis do pliku idzie przez xlsx_patch.py (surowa podmiana komórek w XML), nie przez
openpyxl.Workbook.save() - patrz komentarz w xlsx_patch.py po co (openpyxl przy pełnym
zapisie potrafi popsuć externalLinks/sharedStrings/printerSettings w prawdziwych plikach).
openpyxl jest tu używany tylko do czytania (nagłówki, dopasowanie PPE, sprawdzenie "czy puste").
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

from app import xlsx_patch
from app.excel_reader import (
    FIRST_DATA_ROW,
    SHEET_NAME,
    find_block_starts,
    find_identity_headers,
    find_name_column,
    norm,
    read_block_columns,
)
from app.faktura_reader import (
    KAT_DANE_NIEROZPOZNANE,
    KAT_JUZ_WYPELNIONE,
    KAT_NIEZGODNOSC_SUMY,
    KAT_PPE_NIEZNALEZIONE,
    KAT_SUKCES,
    PozycjaFaktury,
    wczytaj_fakture_dystrybucyjna,
)

MONTH_LABEL_ROW = 1
TOLERANCJA_SUMY_KWH = 0.5

MIESIACE = {
    "STYCZEŃ": 1,
    "LUTY": 2,
    "MARZEC": 3,
    "KWIECIEŃ": 4,
    "MAJ": 5,
    "CZERWIEC": 6,
    "LIPIEC": 7,
    "SIERPIEŃ": 8,
    "WRZESIEŃ": 9,
    "PAŹDZIERNIK": 10,
    "LISTOPAD": 11,
    "GRUDZIEŃ": 12,
}


@dataclass
class WynikWpisu:
    kategoria: str
    ppe: str | None
    wiersz: int | None
    nazwa_obiektu: str | None
    miesiac: int
    okres_od: date | None = None
    okres_do: date | None = None
    opis: str | None = None
    # Nazwa arkusza, do którego trafił ten wpis (Małe odbiory / Duże odbiory / fotowoltaika energia
    # oddana) - do rozróżnienia w UI, kiedy ten sam obiekt/PPE/okres dostaje osobny wpis w dwóch
    # arkuszach naraz z jednej faktury (patrz widok_uzupelnij_excel.py). Celowo POZA kluczem dedupu
    # w import_faktur_polaczony.polacz_i_wyczysc_szum - tylko informacyjne, nie wpływa na logikę.
    arkusz: str | None = None
    # Wpisane (albo dla P-S3 przy 2-strefowej: pominięte) wartości per pole - "P-S1"/"P-S2"/"P-S3"
    # (gdy dotyczy)/"Razem"/"Dystrybucja" -> liczba. Tylko do pokazania w UI (widok_uzupelnij_excel.py,
    # WierszObiektu) - "Razem" to SUMA stref, nie odczyt z arkusza wprost (tam to formuła, patrz
    # xlsx_patch.py), więc może się różnić o grosze od tego co Excel finalnie przeliczy przy
    # zaokrągleniach - to i tak tylko podgląd, nie źródło prawdy (tym jest sam plik Excela).
    wartosci: dict[str, float] | None = None

    @property
    def zapisano(self) -> bool:
        return self.kategoria == KAT_SUKCES


def _bloki_po_miesiacu(ws, kolumny_startowe: list[int]) -> dict[int, dict[str, int]]:
    """Numer miesiąca (1-12) -> mapowanie pole->kolumna, po etykiecie miesiąca w wierszu 1."""
    wynik: dict[int, dict[str, int]] = {}
    for start in kolumny_startowe:
        etykieta = norm(ws.cell(row=MONTH_LABEL_ROW, column=start).value).upper()
        miesiac = MIESIACE.get(etykieta)
        if miesiac is not None:
            wynik[miesiac] = read_block_columns(ws, start)
    return wynik


def _wiersze_po_ppe(ws, name_col: int, ppe_col: int | None) -> dict[str, tuple[int, str]]:
    """Kod PPE -> (numer wiersza, nazwa obiektu), dla wierszy które mają Kod PPE."""
    if ppe_col is None:
        return {}
    wynik: dict[str, tuple[int, str]] = {}
    for row in range(FIRST_DATA_ROW, ws.max_row + 1):
        nazwa = ws.cell(row=row, column=name_col).value
        if nazwa is None or norm(nazwa) == "":
            continue
        ppe = norm(ws.cell(row=row, column=ppe_col).value)
        if ppe:
            wynik[ppe] = (row, norm(nazwa))
    return wynik


def _komorki_puste(ws, wiersz: int, kolumny: dict[str, int]) -> bool:
    for pole in ("data_raw", "zuzycie_szczytowa", "zuzycie_pozaszczytowa", "oplata_dystrybucja"):
        col = kolumny.get(pole)
        if col is not None and norm(ws.cell(row=wiersz, column=col).value) != "":
            return False
    return True


def _data_do_zapisu(pozycja: PozycjaFaktury) -> str:
    return (
        f"{pozycja.okres_od.day:02d}.{pozycja.okres_od.month:02d}"
        f"-{pozycja.okres_do.day:02d}.{pozycja.okres_do.month:02d}"
    )


def _referencja(wiersz: int, kolumna: int) -> str:
    return f"{get_column_letter(kolumna)}{wiersz}"


def _wpisz_pozycje(
    ws,
    pozycje: list[PozycjaFaktury],
    bloki_po_miesiacu: dict[int, dict[str, int]],
    wiersze_po_ppe: dict[str, tuple[int, str]],
    nadpisuj: bool,
) -> tuple[list[WynikWpisu], dict[str, str | float]]:
    wyniki: list[WynikWpisu] = []
    zapisy: dict[str, str | float] = {}
    for pozycja in pozycje:
        miesiac = pozycja.okres_do.month
        trafienie = wiersze_po_ppe.get(pozycja.ppe)
        if trafienie is None:
            wyniki.append(
                WynikWpisu(
                    kategoria=KAT_PPE_NIEZNALEZIONE,
                    ppe=pozycja.ppe,
                    wiersz=None,
                    nazwa_obiektu=None,
                    miesiac=miesiac,
                    okres_od=pozycja.okres_od,
                    okres_do=pozycja.okres_do,
                    opis="brak obiektu z tym numerem PPE w arkuszu",
                    arkusz=SHEET_NAME,
                )
            )
            continue
        wiersz, nazwa_obiektu = trafienie

        kolumny = bloki_po_miesiacu.get(miesiac)
        if kolumny is None:
            wyniki.append(
                WynikWpisu(
                    kategoria=KAT_DANE_NIEROZPOZNANE,
                    ppe=pozycja.ppe,
                    wiersz=wiersz,
                    nazwa_obiektu=nazwa_obiektu,
                    miesiac=miesiac,
                    okres_od=pozycja.okres_od,
                    okres_do=pozycja.okres_do,
                    opis="nie znaleziono bloku dla tego miesiąca w arkuszu",
                    arkusz=SHEET_NAME,
                )
            )
            continue

        if not nadpisuj and not _komorki_puste(ws, wiersz, kolumny):
            wyniki.append(
                WynikWpisu(
                    kategoria=KAT_JUZ_WYPELNIONE,
                    ppe=pozycja.ppe,
                    wiersz=wiersz,
                    nazwa_obiektu=nazwa_obiektu,
                    miesiac=miesiac,
                    okres_od=pozycja.okres_od,
                    okres_do=pozycja.okres_do,
                    opis="komórki już wypełnione (tryb bez nadpisywania)",
                    arkusz=SHEET_NAME,
                )
            )
            continue

        zapisy[_referencja(wiersz, kolumny["data_raw"])] = _data_do_zapisu(pozycja)
        zapisy[_referencja(wiersz, kolumny["zuzycie_szczytowa"])] = pozycja.zuzycie_szczytowa
        # Licznik jednostrefowy/calodobowy: cala ilosc juz jest w zuzycie_szczytowa (patrz
        # faktura_reader.py) - P-S2 celowo zostaje puste, nie piszemy tam literalnego zera.
        if not pozycja.jednostrefowa:
            zapisy[_referencja(wiersz, kolumny["zuzycie_pozaszczytowa"])] = pozycja.zuzycie_pozaszczytowa
        zapisy[_referencja(wiersz, kolumny["oplata_dystrybucja"])] = pozycja.oplata_dystrybucja

        wartosci: dict[str, float] = {"P-S1": pozycja.zuzycie_szczytowa}
        if not pozycja.jednostrefowa:
            wartosci["P-S2"] = pozycja.zuzycie_pozaszczytowa
        wartosci["Razem"] = pozycja.zuzycie_razem
        wartosci["Dystrybucja"] = pozycja.oplata_dystrybucja

        # Niezgodność sumy NIE blokuje zapisu (z decyzji użytkownika, ten sam mechanizm co
        # w Duże odbiory/import_faktur_duze.py) - dane i tak biorą się z faktury, a niezgodność
        # (np. błąd po stronie ENEA) jest tylko ostrzeżeniem do ręcznej weryfikacji.
        suma = pozycja.zuzycie_szczytowa + pozycja.zuzycie_pozaszczytowa
        if abs(suma - pozycja.zuzycie_razem) > TOLERANCJA_SUMY_KWH:
            wyniki.append(
                WynikWpisu(
                    kategoria=KAT_NIEZGODNOSC_SUMY,
                    ppe=pozycja.ppe,
                    wiersz=wiersz,
                    nazwa_obiektu=nazwa_obiektu,
                    miesiac=miesiac,
                    okres_od=pozycja.okres_od,
                    okres_do=pozycja.okres_do,
                    opis=(
                        f"P-S1+P-S2 ({suma:g}) nie zgadza się z zużyciem na fakturze ({pozycja.zuzycie_razem:g})"
                    ),
                    arkusz=SHEET_NAME,
                )
            )
            continue

        wyniki.append(
            WynikWpisu(
                kategoria=KAT_SUKCES,
                ppe=pozycja.ppe,
                wiersz=wiersz,
                nazwa_obiektu=nazwa_obiektu,
                miesiac=miesiac,
                okres_od=pozycja.okres_od,
                okres_do=pozycja.okres_do,
                arkusz=SHEET_NAME,
                wartosci=wartosci,
            )
        )
    return wyniki, zapisy


def wpisz_fakture_do_arkusza(
    sciezka_excel: str | Path, sciezka_faktury: str | Path, nadpisuj: bool
) -> list[WynikWpisu]:
    """Parsuje jedną fakturę dystrybucyjną PDF i wpisuje jej pozycje do arkusza 'Małe odbiory',
    dopasowując po Kodzie PPE. Zapisuje plik Excela na miejscu. Zwraca wynik per pozycja faktury
    (co się stało - zapisano / pominięto i dlaczego), do pokazania użytkownikowi.
    """
    wynik_faktury = wczytaj_fakture_dystrybucyjna(sciezka_faktury)

    # Tylko do czytania - o zapisie patrz komentarz na górze pliku i w xlsx_patch.py.
    wb = openpyxl.load_workbook(sciezka_excel, data_only=False)
    ws = wb[SHEET_NAME]

    name_col = find_name_column(ws)
    kolumny_startowe = find_block_starts(ws)
    pierwszy_blok = min(kolumny_startowe) if kolumny_startowe else ws.max_column + 1
    identity_headers = find_identity_headers(ws, pierwszy_blok)

    bloki_po_miesiacu = _bloki_po_miesiacu(ws, kolumny_startowe)
    wiersze_po_ppe = _wiersze_po_ppe(ws, name_col, identity_headers.get("Kod PPE"))

    wyniki, zapisy = _wpisz_pozycje(ws, wynik_faktury.pozycje, bloki_po_miesiacu, wiersze_po_ppe, nadpisuj)

    for sekcja in wynik_faktury.pominiete:
        wyniki.append(
            WynikWpisu(
                kategoria=sekcja.kategoria,
                ppe=sekcja.ppe,
                wiersz=None,
                nazwa_obiektu=None,
                miesiac=sekcja.okres_do.month if sekcja.okres_do else 0,
                okres_od=sekcja.okres_od,
                okres_do=sekcja.okres_do,
                opis=sekcja.opis,
                arkusz=SHEET_NAME,
            )
        )

    if zapisy:
        xlsx_patch.wpisz_wartosci(sciezka_excel, SHEET_NAME, zapisy)
        # "Razem zużycie"/"Razem koszty" to formuły (=Q+R, =T+U) - nie przeliczamy ich sami,
        # więc bez tego ich zapisana wartość zostałaby stara/zerowa aż do otwarcia w Excelu.
        xlsx_patch.wymus_przeliczenie_formul(sciezka_excel)

    return wyniki
