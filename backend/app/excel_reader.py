"""Parser arkusza 'Małe odbiory' z tabelki rozliczeniowej."""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

import openpyxl

SHEET_NAME = "Małe odbiory"
HEADER_ROW = 2
FIRST_DATA_ROW = 3
NAME_COLUMN_HEADER = "Ulica/ Nazwa własna"

# Nagłówek w bloku miesięcznym -> nazwa pola w naszym modelu.
# Nagłówki spoza tego słownika (np. "Moc pobrana" w bloku Marca) są ignorowane.
FIELD_ALIASES = {
    "Data": "data_raw",
    "P-S1": "zuzycie_szczytowa",
    "P-S2": "zuzycie_pozaszczytowa",
    "Razem zużycie": "zuzycie_razem",
    "Opłata netto dystrybucja": "oplata_dystrybucja",
    "Opłata netto energia el.": "oplata_energia",
    "Razem koszty": "koszty_razem",
}
BLOCK_END_HEADER = "Razem koszty"


def norm(value) -> str:
    return " ".join(str(value).split()) if value is not None else ""


def _komorka(kolumna: int, wiersz: int) -> str:
    return f"{openpyxl.utils.get_column_letter(kolumna)}{wiersz}"


@dataclass
class Okres:
    wiersz: int
    kolumny: dict[str, int]  # nazwa pola -> numer kolumny w arkuszu (do wyliczenia adresu komórki)
    data_raw: str
    data_od: date
    data_do: date
    zuzycie_szczytowa: float
    zuzycie_pozaszczytowa: float
    zuzycie_razem: float
    oplata_dystrybucja: float
    oplata_energia: float
    koszty_razem: float

    def komorka(self, pole: str) -> str:
        return _komorka(self.kolumny[pole], self.wiersz)


@dataclass
class ProblemDanych:
    wiersz: int
    kolumna: int
    tresc: str
    opis: str


@dataclass
class Obiekt:
    wiersz: int
    nazwa: str
    identyfikacja: dict[str, str] = field(default_factory=dict)
    okresy: list[Okres] = field(default_factory=list)
    problemy: list[ProblemDanych] = field(default_factory=list)


def find_block_starts(ws) -> list[int]:
    """Kolumny, w których zaczyna się blok miesięczny (nagłówek 'Data'), z pominięciem ukrytych kolumn."""
    starts = []
    for col in range(1, ws.max_column + 1):
        letter = openpyxl.utils.get_column_letter(col)
        dim = ws.column_dimensions.get(letter)
        if dim is not None and dim.hidden:
            continue
        if norm(ws.cell(row=HEADER_ROW, column=col).value) == "Data":
            starts.append(col)
    return starts


def read_block_columns(ws, start_col: int) -> dict[str, int]:
    """Mapuje nazwę pola -> numer kolumny dla jednego bloku miesięcznego, zaczynając od kolumny 'Data'."""
    columns: dict[str, int] = {}
    col = start_col
    guard = start_col + 15
    while col <= guard:
        header = norm(ws.cell(row=HEADER_ROW, column=col).value)
        alias = FIELD_ALIASES.get(header)
        if alias:
            columns[alias] = col
        if header == BLOCK_END_HEADER:
            break
        col += 1
    return columns


def find_name_column(ws) -> int:
    for col in range(1, ws.max_column + 1):
        if norm(ws.cell(row=HEADER_ROW, column=col).value) == NAME_COLUMN_HEADER:
            return col
    raise ValueError(f"Nie znaleziono kolumny nagłówka {NAME_COLUMN_HEADER!r} w wierszu {HEADER_ROW}")


def find_identity_headers(ws, ostatnia_kolumna: int) -> dict[str, int]:
    """Nagłówek (znormalizowany) -> numer kolumny, dla kolumn tożsamości obiektu (przed pierwszym blokiem
    miesięcznym) — niezależnie od tego, czy kolumna jest ukryta, bo to dane do okna szczegółów obiektu."""
    naglowki: dict[str, int] = {}
    for col in range(1, ostatnia_kolumna):
        header = norm(ws.cell(row=HEADER_ROW, column=col).value)
        if header:
            naglowki[header] = col
    return naglowki


class NieprawidlowyZakresDat(ValueError):
    pass


def _parse_data_range(raw: str, rok: int) -> tuple[date, date]:
    """'01.01-28.02' -> (date(rok,1,1), date(rok,2,28)). Faktura jest zapisana pod miesiącem daty końcowej,
    więc jeśli miesiąc końcowy < miesiąc początkowy, okres przechodzi przez Nowy Rok."""
    parts = raw.split("-")
    if len(parts) != 2:
        raise NieprawidlowyZakresDat(f"oczekiwano formatu 'dd.mm-dd.mm', otrzymano {raw!r}")
    od_str, do_str = parts
    try:
        od_dzien, od_miesiac = (int(x) for x in od_str.split("."))
        do_dzien, do_miesiac = (int(x) for x in do_str.split("."))
        rok_od = rok - 1 if do_miesiac < od_miesiac else rok
        return date(rok_od, od_miesiac, od_dzien), date(rok, do_miesiac, do_dzien)
    except ValueError as exc:
        raise NieprawidlowyZakresDat(f"nie udało się rozpoznać {raw!r}: {exc}") from None


def _to_float(value) -> float:
    if value is None or value == "":
        return 0.0
    return float(value)


def parse_workbook(path: str | Path, rok: int) -> list[Obiekt]:
    """Parsuje arkusz 'Małe odbiory' do listy obiektów z ich okresami rozliczeniowymi.

    `rok` to rok kalendarzowy, którego dotyczy cały skoroszyt (np. 2026) — daty w kolumnie
    'Data' nie zawierają roku, więc trzeba go podać z zewnątrz (nazwa pliku / wybór użytkownika).
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[SHEET_NAME]

    name_col = find_name_column(ws)
    block_starts = find_block_starts(ws)
    blocks = [read_block_columns(ws, start) for start in block_starts]
    pierwszy_blok = min(block_starts) if block_starts else ws.max_column + 1
    identity_headers = find_identity_headers(ws, pierwszy_blok)

    obiekty: list[Obiekt] = []
    for row in range(FIRST_DATA_ROW, ws.max_row + 1):
        nazwa = ws.cell(row=row, column=name_col).value
        if nazwa is None or norm(nazwa) == "":
            continue
        nazwa = norm(nazwa)

        identyfikacja: dict[str, str] = {}
        for header, col in identity_headers.items():
            if header == NAME_COLUMN_HEADER:
                continue
            wartosc = ws.cell(row=row, column=col).value
            if wartosc is not None and norm(wartosc) != "":
                identyfikacja[header] = norm(wartosc)

        okresy: list[Okres] = []
        problemy: list[ProblemDanych] = []
        for block in blocks:
            data_col = block.get("data_raw")
            if data_col is None:
                continue
            data_raw = ws.cell(row=row, column=data_col).value
            if data_raw is None or norm(data_raw) == "":
                continue
            data_raw = norm(data_raw)
            try:
                data_od, data_do = _parse_data_range(data_raw, rok)
            except NieprawidlowyZakresDat as exc:
                problemy.append(ProblemDanych(wiersz=row, kolumna=data_col, tresc=data_raw, opis=str(exc)))
                continue
            okresy.append(
                Okres(
                    wiersz=row,
                    kolumny=block,
                    data_raw=data_raw,
                    data_od=data_od,
                    data_do=data_do,
                    zuzycie_szczytowa=_to_float(ws.cell(row=row, column=block["zuzycie_szczytowa"]).value),
                    zuzycie_pozaszczytowa=_to_float(ws.cell(row=row, column=block["zuzycie_pozaszczytowa"]).value),
                    zuzycie_razem=_to_float(ws.cell(row=row, column=block["zuzycie_razem"]).value),
                    oplata_dystrybucja=_to_float(ws.cell(row=row, column=block["oplata_dystrybucja"]).value),
                    oplata_energia=_to_float(ws.cell(row=row, column=block["oplata_energia"]).value),
                    koszty_razem=_to_float(ws.cell(row=row, column=block["koszty_razem"]).value),
                )
            )

        okresy.sort(key=lambda o: o.data_od)
        obiekty.append(Obiekt(wiersz=row, nazwa=nazwa, identyfikacja=identyfikacja, okresy=okresy, problemy=problemy))

    return obiekty
