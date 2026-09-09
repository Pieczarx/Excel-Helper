"""Odczyt struktury arkusza 'Duże odbiory' - mirror excel_reader.py (Małe odbiory), ale z realnymi
różnicami układu:
- nagłówek jest w wierszu 1 (nie 2), dane od wiersza 2,
- brak wiersza z nazwami miesięcy i brak kolumny 'Data' per blok - bloki miesięczne rozpoznajemy
  PO KOLEJNOŚCI wystąpień nagłówka 'Moc pobrana' (1. = styczeń, ..., 12. = grudzień), nie po treści,
- blok miesięczny ma 16 kolumn (więcej pól: moc pobrana, energia bierna per strefa, opłaty za
  ponadumowny pobór) zamiast 7-8 w Małe odbiory.

Na razie obsługuje tylko odczyt kolumn potrzebny do WPISYWANIA faktur (import_faktur_duze.py) -
nie ma tu odpowiednika parse_workbook() z excel_reader.py (to zasilałoby silnik reguł w zakładce
Weryfikacja dla Dużych odbiorów - poza obecnym zakresem)."""
from __future__ import annotations

import openpyxl

from app.excel_reader import norm

SHEET_NAME_DUZE = "Duże odbiory"
HEADER_ROW_DUZE = 1
FIRST_DATA_ROW_DUZE = 2
NAME_COLUMN_HEADER_DUZE = "Ulica/ Nazwa własna"
BLOCK_START_HEADER_DUZE = "Moc pobrana"
BLOCK_END_HEADER_DUZE = "Razem"

# Nagłówek w bloku miesięcznym -> nazwa pola.
FIELD_ALIASES_DUZE = {
    "Moc pobrana": "moc_pobrana",
    "P-S1": "p_s1",
    "P-S2": "p_s2",
    "P-S3": "p_s3",
    "Q+ S1": "q_plus_s1",
    "Q+ S2": "q_plus_s2",
    "Q+ S3": "q_plus_s3",
    "Q- S1": "q_minus_s1",
    "Q- S2": "q_minus_s2",
    "Q- S3": "q_minus_s3",
    "Q+ (zł)": "q_plus_zl",
    "Q- (zł)": "q_minus_zl",
    "Razem zużycie": "zuzycie_razem",
    "Opłata netto dystrybucja": "oplata_dystrybucja",
    "Opłata netto energia el.": "oplata_energia",
    "Razem": "razem",
}


def find_block_starts_duze(ws) -> list[int]:
    """Kolumny, w których zaczyna się blok miesięczny (nagłówek 'Moc pobrana' w wierszu 1), z
    pominięciem ukrytych kolumn - w kolejności występowania = styczeń..grudzień."""
    starts = []
    for col in range(1, ws.max_column + 1):
        letter = openpyxl.utils.get_column_letter(col)
        dim = ws.column_dimensions.get(letter)
        if dim is not None and dim.hidden:
            continue
        if norm(ws.cell(row=HEADER_ROW_DUZE, column=col).value) == BLOCK_START_HEADER_DUZE:
            starts.append(col)
    return starts


def read_block_columns_duze(ws, start_col: int) -> dict[str, int]:
    """Mapuje nazwę pola -> numer kolumny dla jednego bloku miesięcznego (16 kolumn), zaczynając
    od kolumny 'Moc pobrana'."""
    columns: dict[str, int] = {}
    col = start_col
    guard = start_col + 15
    while col <= guard:
        header = norm(ws.cell(row=HEADER_ROW_DUZE, column=col).value)
        alias = FIELD_ALIASES_DUZE.get(header)
        if alias:
            columns[alias] = col
        if header == BLOCK_END_HEADER_DUZE:
            break
        col += 1
    return columns


def find_name_column_duze(ws) -> int:
    for col in range(1, ws.max_column + 1):
        if norm(ws.cell(row=HEADER_ROW_DUZE, column=col).value) == NAME_COLUMN_HEADER_DUZE:
            return col
    raise ValueError(f"Nie znaleziono kolumny nagłówka {NAME_COLUMN_HEADER_DUZE!r} w wierszu {HEADER_ROW_DUZE}")


def find_identity_headers_duze(ws, ostatnia_kolumna: int) -> dict[str, int]:
    """Nagłówek (znormalizowany) -> numer kolumny, dla kolumn tożsamości obiektu (przed 1. blokiem
    miesięcznym) - niezależnie od tego, czy kolumna jest ukryta."""
    naglowki: dict[str, int] = {}
    for col in range(1, ostatnia_kolumna):
        header = norm(ws.cell(row=HEADER_ROW_DUZE, column=col).value)
        if header:
            naglowki[header] = col
    return naglowki
